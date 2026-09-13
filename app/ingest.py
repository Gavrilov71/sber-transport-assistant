import asyncio
import hashlib
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from .config import get_settings


HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) SBER-Hackathon-RAG/1.0"
}


def clean_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_html(content: bytes) -> str:
    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer"]):
        tag.decompose()
    return clean_text(soup.get_text("\n"))


def extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return clean_text("\n".join(lines))


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    return clean_text("\n".join((page.extract_text() or "") for page in reader.pages))


def extract_text(content: bytes, kind: str) -> str:
    kind = kind.lower()
    if kind == "docx":
        return extract_docx(content)
    if kind == "pdf":
        return extract_pdf(content)
    return extract_html(content)


def chunk_text(text: str, chunk_size: int = 1500, overlap: int = 220) -> list[str]:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n{paragraph}".strip()
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        if current:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            current = f"{tail}\n{paragraph}".strip()
        else:
            for start in range(0, len(paragraph), max(1, chunk_size - overlap)):
                chunks.append(paragraph[start:start + chunk_size])
            current = ""
    if current:
        chunks.append(current)
    return chunks


def safe_filename(source_id: str, kind: str, url: str) -> str:
    suffix = {"html": ".html", "pdf": ".pdf", "docx": ".docx", "doc": ".doc"}.get(kind.lower())
    if not suffix:
        suffix = Path(urlparse(url).path).suffix or ".bin"
    return f"{re.sub(r'[^a-zA-Z0-9._-]+', '_', source_id)}{suffix}"


def discover_documents(base_url: str, content: bytes) -> list[dict]:
    soup = BeautifulSoup(content, "html.parser")
    docs: list[dict] = []
    seen: set[str] = set()
    for link in soup.find_all("a", href=True):
        href = str(link["href"]).strip()
        absolute = urljoin(base_url, href)
        path = urlparse(absolute).path.lower()
        if not path.endswith((".pdf", ".docx", ".doc")):
            continue
        label = " ".join(link.stripped_strings).strip() or Path(path).name
        haystack = f"{label} {absolute}".lower()
        transport_terms = (
            "транспорт", "тройк", "проезд", "льгот", "маршрут", "карт", "асоп",
            "билет", "пассажир", "перевоз", "измен", "асоп", "ткп",
            "83", "59", "412", "3661"
        )
        if not any(term in haystack for term in transport_terms):
            continue
        if absolute in seen:
            continue
        seen.add(absolute)
        if path.endswith(".pdf"):
            kind = "pdf"
        elif path.endswith(".docx"):
            kind = "docx"
        else:
            kind = "doc"
        digest = hashlib.sha1(absolute.encode("utf-8")).hexdigest()[:10]
        docs.append({
            "id": f"oeirc-discovered-{digest}",
            "title": f"ОЕИРЦ — {label}",
            "url": absolute,
            "kind": kind,
            "authority": "АО ОЕИРЦ",
            "category": "нормативные_документы",
            "priority": 68,
            "current": True,
            "enabled": True,
            "discovered": True,
        })
    return docs


def discover_transport_documents(content: bytes, base_url: str, parent: dict) -> list[dict]:
    """Backwards-compatible, testable entry point with parent metadata."""
    docs = discover_documents(base_url, content)
    for doc in docs:
        doc["authority"] = parent.get("authority", doc["authority"])
        doc["priority"] = min(int(parent.get("priority", 75)), 68)
    return docs


async def fetch_source(client: httpx.AsyncClient, source: dict, raw_dir: Path) -> tuple[dict, list[dict]]:
    response = await client.get(source["url"], headers=HEADERS, follow_redirects=True)
    response.raise_for_status()

    raw_dir.mkdir(parents=True, exist_ok=True)
    raw_path = raw_dir / safe_filename(source["id"], source.get("kind", "html"), str(response.url))
    raw_path.write_bytes(response.content)

    kind = source.get("kind", "html")
    if kind == "doc":
        text = ""
    else:
        text = extract_text(response.content, kind)

    discovered = []
    if source.get("discover_documents") and kind == "html":
        discovered = discover_transport_documents(response.content, str(response.url), source)

    loaded = {
        **source,
        "resolved_url": str(response.url),
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "http_status": response.status_code,
        "raw_file": str(raw_path),
        "text": text,
    }
    return loaded, discovered


async def build_chunks() -> list[dict]:
    settings = get_settings()
    sources = json.loads(settings.sources_path.read_text(encoding="utf-8"))
    enabled = [s for s in sources if s.get("enabled", True) and s.get("current", True)]
    results: list[dict] = []
    snapshots: list[dict] = []
    raw_dir = settings.chunks_path.parent / "raw_sources"

    async with httpx.AsyncClient(timeout=45, verify=False, follow_redirects=True) as client:
        queue = list(enabled)
        seen_ids = {s["id"] for s in queue}
        index = 0
        while index < len(queue):
            source = queue[index]
            index += 1
            try:
                loaded, discovered = await fetch_source(client, source, raw_dir)
                snapshots.append(loaded)
                for doc in discovered:
                    if doc["id"] not in seen_ids:
                        seen_ids.add(doc["id"])
                        queue.append(doc)

                if not loaded["text"]:
                    print(f"[INFO] {source['title']}: downloaded only ({source.get('kind')})")
                    continue

                parts = chunk_text(loaded["text"])
                for part_index, text in enumerate(parts):
                    results.append({
                        "id": f"{source['id']}:{part_index}",
                        "source_id": source["id"],
                        "title": source["title"],
                        "url": source["url"],
                        "resolved_url": loaded.get("resolved_url"),
                        "authority": source.get("authority"),
                        "category": source.get("category"),
                        "priority": source.get("priority", 50),
                        "current": source.get("current", True),
                        "fetched_at": loaded.get("fetched_at"),
                        "text": text,
                    })
                print(f"[OK] {source['title']}: {len(parts)} chunks")
            except Exception as exc:
                print(f"[WARN] {source['title']}: {exc}")

    if not results:
        raise RuntimeError("No RAG chunks were built. Check source availability/network.")

    settings.chunks_path.write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    snapshot_path = settings.chunks_path.parent / "source_snapshots.json"
    snapshot_path.write_text(
        json.dumps(snapshots, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Saved {len(results)} chunks to {settings.chunks_path}")
    print(f"Saved {len(snapshots)} source snapshots to {snapshot_path}")
    print(f"Raw source copies: {raw_dir}")
    return results


async def main() -> None:
    chunks = await build_chunks()
    print(f"Agentic text-search corpus is ready: {len(chunks)} chunks")


if __name__ == "__main__":
    asyncio.run(main())
